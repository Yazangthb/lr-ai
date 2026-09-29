import json
import os

import openpyxl
from conftest import paper

from lrai import layout
from lrai.export_gsheets import build_plan, write_plan
from lrai.export_xlsx import write_xlsx

CONFIG = """
taxonomy:
  - name: Forecasting
    rq: RQ1
    children: [Citation counts, Topics]
  - name: Tools
"""


def sample_papers():
    def p(i, **kw):
        return paper(f"Paper number {i} on citation forecasting", doi=f"10.1/{i}", year=2020 + i % 3,
                     citation_count=i, venue="Scientometrics", **kw)
    return [
        p(1, status="core", status_reason="method", first_level="Forecasting", second_level="Citation counts",
          priority="high", summary="S1"),
        p(2, status="skim", status_reason="related", first_level="Forecasting", second_level="Topics"),
        p(3, status="core", first_level="Tools"),
        p(4, status="exclude", first_level="Tools"),  # not included
        p(5, status="core", first_level="Brand new category"),
        paper("=Formula-looking title that must stay text", doi="10.1/6", status="skim"),
    ]


def build(make_run):
    run = make_run(CONFIG)
    stats = {"identified": {"seed": 3}, "identified_total": 3, "unique": 3, "duplicates": 0, "filtered_out": 0,
             "filter_reasons": {}, "screened": 3, "awaiting_screening": 0, "excluded_at_screening": 1,
             "core": 3, "skim": 2, "enriched": 1}
    return run, layout.build(sample_papers(), run.config, stats)


def test_papers_tab_groups_in_taxonomy_order(make_run):
    _, (papers_tab, overview, method) = build(make_run)
    firsts = [r[0] for r, k in zip(papers_tab.rows, papers_tab.kinds) if k == "group"]
    assert firsts == ["▌ Forecasting — 2 papers", "▌ Tools — 1 paper", "▌ Brand new category — 1 paper",
                      "▌ Uncategorized — 1 paper"]
    subs = [r[0] for r, k in zip(papers_tab.rows, papers_tab.kinds) if k == "subgroup"]
    assert subs == ["▸ Citation counts (1)", "▸ Topics (1)"]
    assert papers_tab.kinds.count("paper") == 5  # the excluded paper is left out
    assert overview.name == "Overview" and method.name == "Method"


def test_xlsx_has_styles_rules_and_validation(make_run, tmp_path):
    _, tabs = build(make_run)
    path = str(tmp_path / "out.xlsx")
    write_xlsx(tabs, path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Papers"]
    assert wb.sheetnames == ["Papers", "Overview", "Method"]
    assert ws.freeze_panes == "B2" and ws.auto_filter.ref.startswith("A1:")
    assert ws["A2"].value.startswith("▌ Forecasting") and ws["A2"].font.bold
    assert len(ws.conditional_formatting) >= 3
    assert len(ws.data_validations.dataValidation) == 2
    titles = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
    assert "=Formula-looking title that must stay text" in titles


def test_gsheets_plan_writes_every_row_once_and_keeps_text_literal(make_run, tmp_path):
    _, tabs = build(make_run)
    files = build_plan(tabs, "SHEET123", papers_sheet_id=0, account="acc1", chunk_chars=300)
    names = [n for n, _, _ in files]
    assert names[0] == "setup" and names[-1] == "rules" and "format" in names
    value_calls = [c for n, _, calls in files if n.startswith("values_") and n != "values_summary" for c in calls]
    written = [row for c in value_calls for d in c["arguments"]["data"] for row in d["values"]]
    assert len(written) == len(tabs[0].rows)  # chunks cover the Papers tab exactly
    assert "'=Formula-looking title that must stay text" in [r[0] for r in written]
    assert all(c["account"] == "acc1" for _, _, calls in files for c in calls)
    for _, _, calls in files:
        assert len(calls) <= 50  # one COMPOSIO_MULTI_EXECUTE_TOOL call per file
    paths = write_plan(files, str(tmp_path / "plan"))
    assert os.path.basename(paths[0]) == "01_setup.json"
    with open(paths[0], encoding="utf-8") as f:
        assert json.load(f)["tools"][0]["tool_slug"] == "GOOGLESHEETS_UPDATE_SHEET_PROPERTIES"


def test_sheet_plan_account_comes_from_flag_or_environment(make_run, monkeypatch):
    from lrai import cli
    run, _ = build(make_run)
    run.save(sample_papers())

    def accounts(*extra):
        assert cli.main(["sheet-plan", "--run", run.path, "--spreadsheet-id", "S", *extra]) == 0
        folder = run.file("export", "gsheets")
        found = set()
        for name in os.listdir(folder):
            with open(os.path.join(folder, name), encoding="utf-8") as f:
                found |= {c.get("account") for c in json.load(f)["tools"]}
        return found

    monkeypatch.delenv("LR_AI_GOOGLE_ACCOUNT", raising=False)
    assert accounts() == {None}
    monkeypatch.setenv("LR_AI_GOOGLE_ACCOUNT", "work")
    assert accounts() == {"work"}
    assert accounts("--account", "home") == {"home"}


def test_gsheets_rules_cover_groups_status_priority_and_reading(make_run):
    _, tabs = build(make_run)
    files = dict((n, calls) for n, _, calls in build_plan(tabs, "S"))
    slugs = [c["tool_slug"] for c in files["rules"]]
    assert slugs.count("GOOGLESHEETS_MUTATE_CONDITIONAL_FORMAT_RULES") == 2 + 3 + 3 + 1
    assert slugs.count("GOOGLESHEETS_SET_DATA_VALIDATION_RULE") == 2
    assert slugs[-1] == "GOOGLESHEETS_SET_BASIC_FILTER"
    formulas = [c["arguments"]["rule"]["booleanRule"]["condition"]["values"][0]["userEnteredValue"]
                for c in files["rules"] if c["tool_slug"].endswith("CONDITIONAL_FORMAT_RULES")]
    assert '=LEFT($A2,1)="▌"' in formulas and '=$D2="studied"' in formulas
