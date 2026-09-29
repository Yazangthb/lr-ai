"""Styled .xlsx export (always produced; also the manual-import route into Google Sheets)."""
from __future__ import annotations

from .layout import COLORS, READING_CHOICES, PRIORITY_CHOICES, Tab


def write_xlsx(tabs: list[Tab], path: str) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        from openpyxl.formatting.rule import FormulaRule
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        raise SystemExit("openpyxl is missing: run `pip install -r requirements.txt`") from None

    def fill(hex_color: str) -> PatternFill:
        return PatternFill("solid", start_color=hex_color.lstrip("#"), end_color=hex_color.lstrip("#"))

    def font(hex_color: str | None, bold: bool = False, size: int | None = None, **kw) -> Font:
        return Font(color=hex_color.lstrip("#") if hex_color else None, bold=bold, size=size, **kw)

    top = Alignment(vertical="top")
    top_wrap = Alignment(vertical="top", wrap_text=True)
    wb = Workbook()
    wb.remove(wb.active)

    for tab in tabs:
        ws = wb.create_sheet(tab.name)
        ncols = len(tab.widths)
        last = get_column_letter(ncols)
        for ci, px in enumerate(tab.widths, 1):
            ws.column_dimensions[get_column_letter(ci)].width = round(px / 7, 1)

        level = 0
        for ri, (kind, row) in enumerate(zip(tab.kinds, tab.rows), 1):
            for ci, v in enumerate(row, 1):
                if isinstance(v, str):
                    v = ILLEGAL_CHARACTERS_RE.sub("", v)
                cell = ws.cell(row=ri, column=ci, value=v if v != "" else None)
                if isinstance(v, str) and v.startswith("="):
                    cell.data_type = "s"  # literal text, not a formula
                url = tab.links.get((ri - 1, ci - 1))
                if url:
                    cell.hyperlink = url
                    cell.font = font("#1155CC", underline="single")
                cell.alignment = top_wrap if (ci - 1) in tab.wrap_cols or kind in ("kv", "header") else top
            if kind in ("header", "section", "group", "subgroup"):
                bg, fg = COLORS["header" if kind == "section" else kind]
                for ci in range(1, ncols + 1):
                    c = ws.cell(row=ri, column=ci)
                    c.fill = fill(bg)
                    c.font = font(fg, bold=True)
            elif kind == "title":
                ws.cell(row=ri, column=1).font = font(None, bold=True, size=14)
            elif kind == "kv":
                ws.cell(row=ri, column=1).font = font(None, bold=True)
                ws.cell(row=ri, column=1).fill = fill(COLORS["key"][0])
            # collapsible outline: categories > sub-categories > papers
            if kind == "group":
                level = 1
            elif kind == "subgroup":
                ws.row_dimensions[ri].outline_level = 1
                level = 2
            elif kind == "paper" and level:
                ws.row_dimensions[ri].outline_level = level

        if tab.name != "Papers":
            continue
        n = len(tab.rows)
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{last}{n}"
        ws.sheet_properties.outlinePr.summaryBelow = False
        keys = [c.key for c in tab.columns]

        def col(key: str) -> str | None:
            return get_column_letter(keys.index(key) + 1) if key in keys else None

        if n < 2:
            continue
        if (sc := col("status")):
            for status in ("core", "skim", "exclude"):
                bg, fg = COLORS[status]
                ws.conditional_formatting.add(f"{sc}2:{sc}{n}", FormulaRule(
                    formula=[f'LEFT(${sc}2,{len(status)})="{status}"'], fill=fill(bg), font=font(fg, bold=True)))
        if (pc := col("priority")):
            for prio in PRIORITY_CHOICES:
                bg, fg = COLORS[prio]
                ws.conditional_formatting.add(f"{pc}2:{pc}{n}", FormulaRule(
                    formula=[f'${pc}2="{prio}"'], fill=fill(bg), font=font(fg, bold=prio == "high")))
            dv = DataValidation(type="list", formula1='"' + ",".join(PRIORITY_CHOICES) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{pc}2:{pc}{n}")
        if (rc := col("reading")):
            skip = {col("status"), col("priority")}
            spans, start = [], None
            for i in range(1, ncols + 2):  # contiguous column spans, skipping status/priority
                keep = i <= ncols and get_column_letter(i) not in skip
                if keep and start is None:
                    start = i
                elif not keep and start is not None:
                    spans.append((start, i - 1))
                    start = None
            ranges = " ".join(f"{get_column_letter(a)}2:{get_column_letter(b)}{n}" for a, b in spans)
            ws.conditional_formatting.add(ranges, FormulaRule(
                formula=[f'${rc}2="studied"'], fill=fill(COLORS["studied"][0])))
            dv = DataValidation(type="list", formula1='"' + ",".join(READING_CHOICES) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{rc}2:{rc}{n}")
    wb.save(path)
